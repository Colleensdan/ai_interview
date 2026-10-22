"""The code x interview workbook: what it reads, and that it reaches SharePoint."""

from __future__ import annotations

import io

import openpyxl

import config
from app import state_sync
from models.base import CodeHit
from pipeline import export, storage
from pipeline.agreement import KappaResult
from test_state_sync import fake_sp  # noqa: F401  (fixture)

MODEL = "azure_openai-test"
DOCS = ["R_a.txt", "R_b.txt"]


def _seed(monkeypatch):
    monkeypatch.setattr(config, "EXCLUDED_CODES", ["worry"])
    monkeypatch.setattr(config, "GT_CODE_ALIASES", {"awareness": "attention"})
    conn = storage.init_db(config.DB_PATH)
    storage.insert_ground_truth_matrix(
        conn, {"trust": {"R_a": 1, "R_b": 0}, "attention": {"R_a": 0, "R_b": 2},
               "worry": {"R_a": 1, "R_b": 1}},
        ["trust", "attention", "worry"], ["R_a", "R_b"])
    # Run A found trust in R_b; run B (newer) found it nowhere. The site shows
    # run B's kappa, so the workbook must show run B's codings.
    storage.insert_count_matrix(conn, "run_a", MODEL, {"trust": {"R_a.txt": 0, "R_b.txt": 3}},
                                ["trust"], DOCS)
    storage.insert_kappa(conn, "run_a", [(MODEL, KappaResult("trust", 0.2, 2, 1, 1))])
    storage.insert_count_matrix(conn, "run_b", MODEL, {"trust": {"R_a.txt": 0, "R_b.txt": 0}},
                                ["trust"], DOCS)
    storage.insert_kappa(conn, "run_b", [(MODEL, KappaResult("trust", 0.0, 2, 1, 0))])
    storage.insert_majority_matrix(conn, "run_b", {"trust": {"R_a.txt": 0, "R_b.txt": 0}},
                                   ["trust"], DOCS)
    storage.insert_kappa(conn, "run_b", [("majority_vote", KappaResult("trust", 0.0, 2, 1, 0))])
    storage.insert_kappa(conn, "run_old", [(MODEL, KappaResult("worry", 0.23, 2, 2, 1))])
    conn.commit()
    return conn


def _sheet(wb, name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return {r[0]: dict(zip(header[1:], r[1:])) for r in rows[1:]}


def test_workbook_reads_each_column_from_the_kappa_run(mem_db, monkeypatch):
    conn = _seed(monkeypatch)
    wb = openpyxl.load_workbook(io.BytesIO(
        export.workbook_bytes(conn, ["trust", "awareness", "worry"])))
    conn.close()

    assert wb.sheetnames == ["Human ground truth", "Majority vote", MODEL, "Notes"]
    human = _sheet(wb, "Human ground truth")
    assert set(human) == {"R_a", "R_b"}
    assert human["R_a"] == {"trust": 1, "awareness": 0}   # renamed code via alias
    assert human["R_b"] == {"trust": 0, "awareness": 2}
    model = _sheet(wb, MODEL)
    assert model["R_b"]["trust"] == 0, "run A's count leaked under run B's kappa"
    assert model["R_b"]["awareness"] is None, "never coded must be blank, not 0"
    assert "worry" not in [c.value for c in wb[MODEL][1]], "withdrawn code exported"
    notes = [r for r in wb["Notes"].iter_rows(values_only=True) if r and r[0] == MODEL]
    assert notes == [(MODEL, "trust", "run_b", 0.0, None)]


def test_push_uploads_the_workbook(fake_sp, monkeypatch):  # noqa: F811
    _seed(monkeypatch).close()
    state_sync.push_state()
    path = state_sync._export_remote(export.WORKBOOK_NAME)
    assert path in fake_sp.files
    wb = openpyxl.load_workbook(io.BytesIO(fake_sp.files[path][0]))
    assert MODEL in wb.sheetnames


def test_coding_results_csv_pairs_quotes_with_the_kappa_run(fake_sp, monkeypatch):  # noqa: F811
    conn = _seed(monkeypatch)
    storage.insert_coding_results(conn, "run_a", MODEL,
                                  [CodeHit("R_b.txt", "trust", "old quote", "r")],
                                  code_names=["trust"])
    conn.commit()
    conn.close()
    state_sync.push_state()
    csv_bytes = fake_sp.files[state_sync._csv_remote("coding_results.csv")][0]
    assert b"old quote" not in csv_bytes
